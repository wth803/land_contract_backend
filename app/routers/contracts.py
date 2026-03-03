# 土地承包明细 API 路由端点
import io
from datetime import datetime
from typing import Optional
from urllib.parse import quote

import openpyxl
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy.orm import Session

from app import crud, models, schemas
from app.auth import get_current_user
from app.database import get_db

router = APIRouter(prefix="/api/contracts", tags=["土地承包明细"])


@router.get("", response_model=schemas.PaginatedResponse, summary="获取土地承包明细列表")
def list_contracts(
    page: int = Query(1, ge=1, description="页码，从1开始"),
    page_size: int = Query(10, ge=1, le=100, description="每页条数"),
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    """获取土地承包明细列表，支持分页，按创建时间倒序排列"""
    total, items = crud.get_contracts(db, page=page, page_size=page_size)
    return schemas.PaginatedResponse(
        total=total, page=page, page_size=page_size, items=items
    )


@router.get("/search", response_model=schemas.PaginatedResponse, summary="模糊搜索土地承包明细")
def search_contracts(
    name: Optional[str] = Query(None, description="承包人姓名（模糊匹配）"),
    land_location: Optional[str] = Query(None, description="地块位置（模糊匹配）"),
    page: int = Query(1, ge=1, description="页码，从1开始"),
    page_size: int = Query(10, ge=1, le=100, description="每页条数"),
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    """通过承包人姓名和地块位置进行模糊搜索，支持分页"""
    total, items = crud.search_contracts(
        db, name=name, land_location=land_location, page=page, page_size=page_size
    )
    return schemas.PaginatedResponse(
        total=total, page=page, page_size=page_size, items=items
    )


@router.post(
    "/export",
    summary="导出土地承包明细到 Excel",
    response_class=StreamingResponse,
)
def export_contracts(
    export_req: schemas.ExportRequest,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    """
    根据所选列和搜索条件导出土地承包明细为 Excel 文件
    """
    # 字段名到中文表头的映射
    column_headers = {
        "name": "承包人姓名",
        "id_card": "身份证号",
        "phone": "联系电话",
        "land_location": "地块位置",
        "area": "承包面积（亩）",
        "year": "承包年份",
        "village": "村别",
        "bank_account": "银行卡号",
        "contractor_code": "承包方编码",
        "plot_code": "地块编码",
        "remark": "备注",
        "created_at": "创建时间",
        "updated_at": "更新时间",
    }

    # 获取匹配记录（不分页）
    items = crud.get_all_contracts(
        db,
        name=export_req.search_name,
        land_location=export_req.search_land_location,
    )

    # 无数据时返回提示
    if not items:
        raise HTTPException(status_code=404, detail="没有符合条件的数据可供导出")

    # 创建 Excel 工作簿
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "土地承包明细"

    # 表头样式：加粗 + 浅绿色背景
    header_font = Font(bold=True)
    header_fill = PatternFill(fill_type="solid", fgColor="C6EFCE")
    header_alignment = Alignment(horizontal="center", vertical="center")

    # 写入表头行，同时初始化各列最大宽度
    headers = [column_headers[col] for col in export_req.columns]
    col_widths = []
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        # 中文字符按2个宽度单位计算，ASCII字符按1个单位
        col_widths.append(sum(2 if ord(c) > 127 else 1 for c in header))

    # 写入数据行，同时追踪各列最大字符宽度
    for row_idx, item in enumerate(items, start=2):
        for col_idx, col_name in enumerate(export_req.columns, start=1):
            value = getattr(item, col_name, None)
            # 将 datetime 对象格式化为字符串
            if isinstance(value, datetime):
                value = value.strftime("%Y-%m-%d %H:%M:%S")
            ws.cell(row=row_idx, column=col_idx, value=value)
            # 更新该列的最大宽度
            if value is not None:
                cell_width = sum(2 if ord(c) > 127 else 1 for c in str(value))
                col_widths[col_idx - 1] = max(col_widths[col_idx - 1], cell_width)

    # 列宽自适应：根据内容最大宽度加边距设置列宽
    for col_idx, width in enumerate(col_widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width + 4

    # 将工作簿写入内存字节流
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    # 生成带时间戳的文件名
    filename = datetime.now().strftime("土地承包明细_%Y%m%d_%H%M%S.xlsx")
    # 对文件名进行 URL 编码以支持中文
    encoded_filename = quote(filename)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"},
    )


@router.post(
    "",
    response_model=schemas.ContractResponse,
    status_code=201,
    summary="创建土地承包明细",
)
def create_contract(
    contract: schemas.ContractCreate,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    """创建新的土地承包明细记录"""
    return crud.create_contract(db, contract)


@router.put("/{contract_id}", response_model=schemas.ContractResponse, summary="更新土地承包明细")
def update_contract(
    contract_id: int,
    contract: schemas.ContractUpdate,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    """更新指定 ID 的土地承包明细（支持部分字段更新）"""
    db_contract = crud.update_contract(db, contract_id, contract)
    if not db_contract:
        raise HTTPException(status_code=404, detail=f"未找到 ID 为 {contract_id} 的土地承包明细")
    return db_contract


@router.delete("/{contract_id}", summary="删除土地承包明细")
def delete_contract(
    contract_id: int,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    """删除指定 ID 的土地承包明细"""
    success = crud.delete_contract(db, contract_id)
    if not success:
        raise HTTPException(status_code=404, detail=f"未找到 ID 为 {contract_id} 的土地承包明细")
    return {"message": f"ID 为 {contract_id} 的土地承包明细已成功删除"}
